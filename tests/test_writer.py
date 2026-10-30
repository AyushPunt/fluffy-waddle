import io
import re

import docx
import pytest

from docagent.compiler import approve
from docagent.errors import PlanError
from docagent.manifest import Anchor, Manifest, Mutation, Region, RegionKind
from docagent.ooxml.package import Package
from docagent.plan import FillPlan, SystemRecord
from docagent.writer import fill
from tests.fixtures import review
from tests.fixtures.csra import minimal


def _doc_xml(output: bytes) -> bytes:
    return Package.from_bytes(output).part("word/document.xml")


def test_output_reads_back_with_every_region_filled(filled):
    d = docx.Document(io.BytesIO(filled.output))
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    assert tables[0][1:] == [
        ["System Name", "LabTrack LIMS"],
        ["System ID", "SYS-0042"],
        ["System Version", "4.2.1"],
        ["Vendor", "[GAP-1: system.vendor is not in the system record — SME input required]"],
    ]
    assert tables[1][1][0] == "REQ-001" and tables[1][2][0] == "REQ-002"
    assert tables[1][2][3] == "[GAP-3: risk rating for the e-signature requirement — SME input required]"
    assert [r[1] for r in tables[2][1:]] == ["☒ Yes   ☐ No", "☐ Yes   ☒ No", "☐ Yes   ☐ No"]
    assert tables[3] == [["Role", "Name", "Signature", "Date"], ["Author", "", "", ""], ["QA Approver", "", "", ""]]
    texts = [p.text for p in d.paragraphs]
    assert "The business owner of LabTrack LIMS is accountable for this assessment." in texts
    assert "LabTrack LIMS records sample receipt, test results & reviews for QC <Building 4>." in texts
    assert d.sections[0].header.paragraphs[0].text == "Risk Assessment SYS-0042 — Document Version 1.0"


def test_report_lists_gaps_blockers_and_human_regions(filled):
    r = filled.report
    assert [g["n"] for g in r["gaps"]] == [1, 2, 3]
    assert {b["item"] for b in r["blocking"]} == {"Data integrity (ALCOA+): unknown", "Backups verified: unknown"}
    assert r["awaiting_human"] == ["approvals"]
    assert r["gate"]["passed"] and r["status"] == "draft_for_review"
    assert r["boilerplate"] == {"csra_purpose": review.entry_sha256(review.PURPOSE)}


def test_parts_without_regions_are_raw_identical(template, filled):
    t, o = Package.from_bytes(template), Package.from_bytes(filled.output)
    changed = [n for n in t.names if t.entry(n).raw_record != o.entry(n).raw_record]
    assert changed == ["word/document.xml", "word/header1.xml"]


def test_checkbox_state_is_flipped_in_place(filled):
    xml = _doc_xml(filled.output)
    states = re.findall(rb'<w14:checked w14:val="(\d)"/>', xml)
    assert states == [b"1", b"1", b"0", b"0"]  # Part 11, Annex 11, GAMP no, ALCOA+ unknown
    defaults = re.findall(rb'<w:default w:val="(\d)"/>', xml)
    assert defaults == [b"0", b"1", b"0"]  # Medium
    assert xml.count("☒".encode()) == 2 + 2  # two sdt boxes ticked, two yes/no glyphs ticked


def test_split_run_placeholder_keeps_every_run(template, filled):
    xml = _doc_xml(filled.output)
    assert (
        b'<w:r w:rsidR="00B11C3D"><w:t xml:space="preserve">4.2.1</w:t></w:r><w:proofErr w:type="spellStart"/>'
        b'<w:r w:rsidR="00C77E21"><w:t xml:space="preserve"></w:t></w:r><w:proofErr w:type="spellEnd"/>'
        b'<w:r w:rsidR="00D88F32"><w:t xml:space="preserve"></w:t></w:r>' in xml
    )


def test_gap_markers_are_highlighted(filled):
    xml = _doc_xml(filled.output)
    assert re.search(rb'<w:highlight w:val="yellow"/></w:rPr><w:t xml:space="preserve">\[GAP-2:', xml)


def test_inline_gap_splits_the_run_and_passes_the_gate(template, manifest):
    rec = SystemRecord(record_id="S", record_version="1", system={"id": "SYS-1", "version": "1"})
    out = fill(template, manifest, rec, review.plan(manifest), review.library())
    d = docx.Document(io.BytesIO(out.output))
    line = next(p.text for p in d.paragraphs if p.text.startswith("The business owner"))
    assert line == "The business owner of [GAP-1: system.name is not in the system record — SME input required] " \
        "is accountable for this assessment."
    assert out.report["gate"]["passed"]


def test_without_a_plan_every_authored_region_is_a_gap_or_blocker(template, manifest):
    out = fill(template, manifest, review.record(), None, review.library())
    by_id = {r["id"]: r for r in out.report["regions"]}
    assert by_id["system_description"]["status"] == "gap"
    assert by_id["requirements_rows"]["status"] == "gap"
    assert by_id["overall_risk"]["status"] == "blocked"
    assert len(by_id["gxp_applicability"]["blocking"]) == 4


def test_output_is_deterministic(template, manifest, filled):
    again = fill(template, manifest, review.record(), review.plan(manifest), review.library())
    assert again.output == filled.output


def _plan(manifest, regions):
    return FillPlan(manifest_sha256=manifest.approval.manifest_sha256, regions=regions)


@pytest.mark.parametrize(
    "regions, message",
    [
        ({"system_name": {"paragraphs": ["x"]}}, "code fills it"),
        ({"approvals": {"paragraphs": ["x"]}}, "code fills it"),
        ({"overall_risk": {"choice": "Extreme"}}, "option set"),
        ({"system_description": {"paragraphs": ["see [GAP-4: x]"]}}, "numbers them"),
        ({"system_description": {"paragraphs": ["{{system.name"]}}, "unbalanced"),
        ({"system_description": {"paragraphs": ["{{shell}}"]}}, "unknown token"),
        ({"system_description": {"paragraphs": ["two\nlines"]}}, "line break"),
        ({"system_description": {"paragraphs": ["☒ audit trail"]}}, "checkbox glyphs"),
        ({"requirements_rows": {"rows": [{"requirement": "x"}]}}, "exactly the columns"),
        ({"gxp_applicability": {"items": {"ISO 13485": "yes"}}}, "not in the template"),
        ({"nope": {}}, "unknown region"),
    ],
)
def test_bad_plans_are_refused(template, manifest, regions, message):
    with pytest.raises(PlanError, match=message):
        fill(template, manifest, review.record(), _plan(manifest, regions), review.library())


def test_plan_for_another_manifest_is_refused(template, manifest):
    plan = FillPlan(manifest_sha256="0" * 64, regions={})
    with pytest.raises(PlanError, match="different manifest"):
        fill(template, manifest, review.record(), plan, review.library())


def test_tampered_boilerplate_is_refused(template, manifest):
    lib = review.library()
    lib.entries["csra_purpose"].paragraphs = ["Something else."]
    with pytest.raises(PlanError, match="does not match its recorded hash"):
        fill(template, manifest, review.record(), None, lib)


def test_special_characters_survive(template, manifest):
    text = 'Ampersand & angle <b> quotes "x" \'y\' emoji 🧪 accents é'
    out = fill(template, manifest, review.record(), _plan(manifest, {"system_description": {"paragraphs": [text]}}),
               review.library())
    assert text in [p.text for p in docx.Document(io.BytesIO(out.output)).paragraphs]


def test_append_mode_keeps_the_prototype_row(template):
    m = review.reviewed_manifest(template)
    m.region("requirements_rows").row_mode = "append"
    m = approve(m, template, "QA")
    out = fill(template, m, review.record(), review.plan(m), review.library())
    rows = docx.Document(io.BytesIO(out.output)).tables[1].rows
    assert [r.cells[0].text for r in rows] == ["Req ID", "", "REQ-001", "REQ-002"]


def test_section_break_moves_to_the_last_new_paragraph():
    body = (
        '<w:p><w:pPr><w:sectPr><w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/></w:sectPr></w:pPr>'
        "<w:r><w:t>[Describe the landscape section]</w:t></w:r></w:p>"
    )
    raw = minimal(body)
    from docagent.compiler import TemplateDoc

    m = Manifest(
        template=TemplateDoc(raw).template_ref(),
        regions=[Region(id="d", kind=RegionKind.PROSE, mutation=Mutation.PARAGRAPHS,
                        anchor=Anchor(text="[Describe the landscape section]"))],
    )
    m = approve(m, raw, "QA")
    plan = _plan(m, {"d": {"paragraphs": ["One.", "Two.", "Three."]}})
    out = fill(raw, m, review.record(), plan)
    xml = _doc_xml(out.output)
    paras = re.findall(rb"<w:p>.*?</w:p>", xml)
    assert [p.count(b"<w:sectPr>") for p in paras] == [0, 0, 1]
    assert out.report["gate"]["passed"]

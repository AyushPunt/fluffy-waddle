"""Drafting pipeline with a scripted model: every bad claim must end up as a gap or unknown."""

import io
from types import SimpleNamespace

import docx
import pytest

from docagent.author import JUDGE_RULES, author
from docagent.errors import ManifestError
from docagent.kb import mask_record_values
from docagent.llm import AnthropicLLM, LLMResult
from docagent.plan import SystemRecord
from docagent.verify import unsourced_terms
from docagent.writer import fill
from tests.fixtures import review
from tests.fixtures.kb import build_kb


def claim(text, chunk_id, quote, topic="t"):
    return {"kind": "claim", "topic": topic, "text": text, "chunk_id": chunk_id, "quote": quote}


def gap(topic):
    return {"kind": "gap", "topic": topic, "text": "", "chunk_id": "", "quote": ""}


DRAFTS = {
    "system_description": {"paragraphs": [{"items": [
        claim("{{system.name}} manages sample login, test result entry and result review for the QC laboratory.",
              "urs-0042-01", "{{system.name}} manages sample login, test result entry and result review for the QC laboratory"),
        claim("It is used by 45 named analysts.", "vsr-0017-01", "ChromaVault CDS is used by 45 named analysts",
              topic="number of users"),
        claim("QC reviewers approve results within 24 hours.", "urs-0042-02", "QC reviewers, who review and approve them",
              topic="review turnaround"),
        claim("The system replaces all paper records.", "urs-0042-01",
              "It replaces paper worksheets for release testing of finished product", topic="paper records"),
        gap("number of named users at go-live"),
        claim("[ORG_1] IT administers the application servers.", "urs-0042-02",
              "[ORG_1] IT administers the application servers", topic="system administration"),
    ]}]},
    "requirements_rows": {"rows": [
        {"cells": {
            "requirement": claim("The system shall record a secure, time-stamped audit trail of every creation, "
                                 "modification and deletion of a result.", "urs-0042-03",
                                 "The system shall record a secure, time-stamped audit trail of every creation, "
                                 "modification and deletion of a result"),
            "gxp_impact": claim("Direct", "ra-0042-01", "Audit trail of result changes: GxP impact Direct"),
            "risk": claim("High", "ra-0042-01", "Audit trail of result changes: GxP impact Direct, risk High"),
        }},
        {"cells": {
            "requirement": claim("The system shall back up its database every day.", "urs-0042-05",
                                 "The system shall back up its database every day"),
            "gxp_impact": gap("GxP impact of daily backup"),
            "risk": gap("risk rating of daily backup"),
        }},
        {"cells": {
            "requirement": claim("Result approval shall use an electronic signature.", "urs-0042-04",
                                 "Approval uses e-signatures", topic="e-signature requirement"),
            "gxp_impact": claim("Direct", "ra-0042-01", "Electronic signature on result approval: GxP impact Direct"),
            "risk": claim("High", "ra-0042-01", "Electronic signature on result approval: GxP impact Direct, risk High"),
        }},
    ]},
    "gxp_applicability": {"items": [
        {"item": "21 CFR Part 11", "decision": "yes", "chunk_id": "vsr-0042-01",
         "quote": "so 21 CFR Part 11 and EU GMP Annex 11 apply"},
        {"item": "EU GMP Annex 11", "decision": "yes", "chunk_id": "vsr-0042-01",
         "quote": "so 21 CFR Part 11 and EU GMP Annex 11 apply"},
        {"item": "GAMP 5 Category 4", "decision": "yes", "chunk_id": "vsr-0042-01",
         "quote": "is classified as GAMP 5 Category 4"},
        {"item": "Data integrity (ALCOA+)", "decision": "yes", "chunk_id": "oq-0042-01",
         "quote": "verified that the audit trail is enabled"},
    ]},
    "data_integrity_controls": {"items": [
        {"item": "Audit trail enabled", "decision": "yes", "chunk_id": "oq-0042-01",
         "quote": "verified that the audit trail is enabled"},
        {"item": "Electronic signatures used", "decision": "yes", "chunk_id": "oq-0042-01",
         "quote": "Electronic signatures are used for result approval"},
        {"item": "Backups verified", "decision": "yes", "chunk_id": "vsr-0017-01",
         "quote": "Backups of ChromaVault CDS are verified monthly"},
    ]},
}


class FakeLLM:
    """Scripted drafter and judge. `reject` lists judge refs to mark not entailed."""

    model = "fake-model"

    def __init__(self, choices=("unknown", "unknown", "unknown"), reject=("p1.4",), refuse=()):
        self.choices = list(choices)
        self.reject = set(reject)
        self.refuse = set(refuse)
        self.calls = []

    def complete_json(self, system, user, schema, effort):
        self.calls.append((system, user, effort))
        if system[0]["text"] == JUDGE_RULES:
            refs = [line.split('"')[1] for line in user.splitlines() if line.startswith("<pair id=")]
            return self._result({"verdicts": [{"id": r, "entailed": r not in self.reject} for r in refs]})
        region = user.split("\n", 1)[0].removeprefix("Region: ")
        if region in self.refuse:
            return LLMResult(None, self.model, "p", "", None, refused=True)
        if region == "overall_risk":
            pick = self.choices.pop(0)
            return self._result({"choice": pick, "chunk_id": "ra-0042-01",
                                 "quote": "Audit trail of result changes: GxP impact Direct, risk High"})
        return self._result(DRAFTS[region])

    def _result(self, data):
        return LLMResult(data, self.model, "p" * 64, "o" * 64, "req_fake")


@pytest.fixture(scope="module")
def drafted(manifest):
    return author(manifest, review.record(), build_kb(), FakeLLM())


def test_supported_claims_pass_and_bad_ones_become_gaps(drafted):
    text = drafted.plan.regions["system_description"]["paragraphs"][0]
    assert text.startswith("{{system.name}} manages sample login")
    assert "45" not in text and "24 hours" not in text and "all paper records" not in text
    assert "{{GAP: number of users (no verified evidence)}}" in text  # other system's chunk
    assert "{{GAP: review turnaround (no verified evidence)}}" in text  # number not in the quote
    assert "{{GAP: paper records (no verified evidence)}}" in text  # judge said not entailed
    assert "{{GAP: number of named users at go-live}}" in text  # gap declared by the model
    assert "{{GAP: system administration (de-identified detail)}}" in text  # surrogate copied into a claim


def test_audit_records_why_each_claim_failed(drafted):
    claims = {c["ref"]: c for r in drafted.audit["regions"] for c in r["claims"]}
    assert claims["p1.1"]["accepted"]
    assert claims["p1.2"]["checks"]["same_system"] is False
    assert claims["p1.3"]["checks"]["facts_sourced"] is False
    assert claims["p1.4"]["checks"]["entailed"] is False
    assert claims["r3.requirement"]["checks"]["quote_in_chunk"] is False
    assert claims["p1.6"]["checks"]["text_allowed"] is False
    assert "vsr-0017-01" not in drafted.audit["evidence_chunk_ids"]
    assert all("quote" not in c and "text" not in c for c in claims.values())  # hashes only


def test_table_rows_keep_verified_cells_only(drafted):
    rows = drafted.plan.regions["requirements_rows"]["rows"]
    assert rows[0] == {
        "requirement": "The system shall record a secure, time-stamped audit trail of every creation, modification "
                       "and deletion of a result.",
        "gxp_impact": "Direct",
        "risk": "High",
    }
    assert rows[1]["risk"] == "{{GAP: risk rating of daily backup}}"
    assert rows[2]["requirement"] == "{{GAP: e-signature requirement (no verified evidence)}}"


def test_checklist_answers_need_verified_citations(drafted):
    gxp = drafted.plan.regions["gxp_applicability"]["items"]
    assert gxp == {"21 CFR Part 11": "yes", "EU GMP Annex 11": "yes", "GAMP 5 Category 4": "yes",
                   "Data integrity (ALCOA+)": "unknown"}
    di = drafted.plan.regions["data_integrity_controls"]["items"]
    assert di == {"Audit trail enabled": "yes", "Electronic signatures used": "yes", "Backups verified": "unknown"}


def test_a_no_decision_may_cite_evidence_that_states_something_else(manifest):
    class NoLLM(FakeLLM):
        def complete_json(self, system, user, schema, effort):
            if user.startswith("Region: gxp_applicability"):
                return self._result({"items": [{"item": "GAMP 5 Category 4", "decision": "no", "chunk_id": "vsr-0042-01",
                                                "quote": "is a configured product and is classified as GAMP 5 Category"}]})
            return super().complete_json(system, user, schema, effort)

    out = author(manifest, review.record(), build_kb(), NoLLM())
    assert out.plan.regions["gxp_applicability"]["items"]["GAMP 5 Category 4"] == "no"


def test_single_choice_needs_agreeing_samples(manifest):
    split = author(manifest, review.record(), build_kb(), FakeLLM(choices=["High", "High", "Medium"]))
    assert split.plan.regions["overall_risk"] == {"choice": None}
    note = next(r["note"] for r in split.audit["regions"] if r["region"] == "overall_risk")
    assert note == "samples disagreed: High, Medium"
    agreed = author(manifest, review.record(), build_kb(), FakeLLM(choices=["High"] * 3))
    assert agreed.plan.regions["overall_risk"] == {"choice": "High"}


def test_drafted_plan_fills_and_passes_the_gate(template, manifest, drafted):
    out = fill(template, manifest, review.record(), drafted.plan, review.library())
    assert out.report["gate"]["passed"]
    d = docx.Document(io.BytesIO(out.output))
    texts = [p.text for p in d.paragraphs]
    assert any(t.startswith("LabTrack LIMS manages sample login") for t in texts)
    reqs = [[c.text for c in r.cells] for r in d.tables[1].rows]
    assert reqs[1][:2] == ["REQ-001", drafted.plan.regions["requirements_rows"]["rows"][0]["requirement"]]
    blocked = {b["item"] for b in out.report["blocking"]}
    assert {"Data integrity (ALCOA+): unknown", "Backups verified: unknown", "no option chosen"} <= blocked


def test_refusal_leaves_the_region_to_the_writer(manifest):
    out = author(manifest, review.record(), build_kb(), FakeLLM(refuse={"system_description"}))
    assert "system_description" not in out.plan.regions
    region = next(r for r in out.audit["regions"] if r["region"] == "system_description")
    assert region["calls"][0]["refused"] is True


def test_no_evidence_means_no_model_calls(manifest):
    rec = SystemRecord(record_id="SYS-9999", record_version="1", system={"name": "Other"})
    llm = FakeLLM()
    out = author(manifest, rec, build_kb(), llm)
    assert llm.calls == [] and out.plan.regions == {}


def test_unapproved_manifest_is_refused(manifest):
    edited = manifest.model_copy(deep=True)
    edited.approval = None
    with pytest.raises(ManifestError):
        author(edited, review.record(), build_kb(), FakeLLM())


def test_record_values_are_masked_before_the_model_sees_them():
    text = "LabTrack LIMS 4.2.1 (SYS-0042) supersedes v3.1; Annex 11 applies; version 4.2.1."
    masked = mask_record_values(text, review.record())
    assert masked == ("{{system.name}} {{system.version}} ({{system.id}}) supersedes v3.1; Annex 11 applies; "
                      "version {{system.version}}.")


def test_unsourced_terms():
    assert unsourced_terms("Backups are verified every 30 days by QA.", "backups are verified") == ["30", "QA"]
    assert unsourced_terms("{{system.name}} is used by Analysts.", "used by analysts") == []


def test_anthropic_adapter_sends_a_pinned_schema_request():
    sent = {}

    def create(**kwargs):
        sent.update(kwargs)
        return SimpleNamespace(stop_reason="end_turn", model="claude-opus-5-5", _request_id="req_1",
                               content=[SimpleNamespace(type="text", text='{"ok": true}')])

    client = SimpleNamespace(messages=SimpleNamespace(create=create))
    llm = AnthropicLLM(client=client)
    result = llm.complete_json([{"type": "text", "text": "rules"}], "task", {"type": "object"}, "high")
    assert result.data == {"ok": True} and result.request_id == "req_1"
    assert sent["model"] == "claude-opus-5-5"
    assert sent["output_config"] == {"effort": "high", "format": {"type": "json_schema", "schema": {"type": "object"}}}
    assert "temperature" not in sent and "thinking" not in sent
